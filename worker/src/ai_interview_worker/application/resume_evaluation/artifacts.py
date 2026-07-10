"""简历评估产物生成。"""

import os
import tempfile
from io import BytesIO
from typing import Any


def render_radar_chart(result_summary: dict[str, Any]) -> bytes:
    """生成雷达图 PNG bytes。"""
    dimensions = result_summary.get("dimensions", [])
    if not dimensions:
        raise ValueError("dimensions are required")
    os.environ.setdefault("MPLCONFIGDIR", os.path.join(tempfile.gettempdir(), "ai-interview-matplotlib"))
    import matplotlib

    matplotlib.use("Agg")
    import matplotlib.pyplot as plt
    import numpy as np

    labels = [_dimension_label(str(item["name"])) for item in dimensions]
    scores = [_score(item) for item in dimensions]
    angles = np.linspace(0, 2 * np.pi, len(labels), endpoint=False).tolist()
    scores_closed = scores + scores[:1]
    angles_closed = angles + angles[:1]

    plt.rcParams["font.sans-serif"] = ["Arial Unicode MS", "Microsoft YaHei", "SimHei", "DejaVu Sans"]
    plt.rcParams["axes.unicode_minus"] = False
    figure, axis = plt.subplots(figsize=(7, 7), subplot_kw={"polar": True})
    figure.patch.set_facecolor("white")
    axis.set_facecolor("white")
    axis.set_ylim(0, 100)
    axis.set_xticks(angles)
    axis.set_xticklabels(labels, fontsize=10)
    axis.set_yticks([20, 40, 60, 80, 100])
    axis.set_yticklabels(["20", "40", "60", "80", "100"], fontsize=8)
    axis.grid(color="#d7dde8", linewidth=0.8)
    axis.plot(angles_closed, scores_closed, color="#2563eb", linewidth=2.5, marker="o")
    axis.fill(angles_closed, scores_closed, color="#2563eb", alpha=0.18)
    axis.set_title("简历综合评价雷达图", fontsize=15, fontweight="bold", pad=24)
    for angle, score in zip(angles, scores, strict=True):
        axis.text(angle, min(score + 6, 105), str(score), ha="center", va="center", fontsize=9, color="#111827")

    buffer = BytesIO()
    figure.savefig(buffer, format="png", dpi=180, bbox_inches="tight")
    plt.close(figure)
    return buffer.getvalue()


def render_evaluation_report(result_summary: dict[str, Any], radar_chart: bytes | None) -> bytes:
    """生成评估报告 DOCX bytes。"""
    if "overall_score" not in result_summary:
        raise ValueError("overall_score is required")
    summary = str(result_summary.get("summary", "")).strip()
    dimensions = result_summary.get("dimensions", [])
    if not isinstance(dimensions, list) or not dimensions:
        raise ValueError("dimensions are required")
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml.ns import qn
    from docx.shared import Inches, Pt, RGBColor

    document = Document()
    normal_style = document.styles["Normal"]
    normal_style.font.name = "微软雅黑"
    normal_style.font.size = Pt(10.5)
    normal_style.element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")

    title = document.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title_run = title.add_run("简历评估报告")
    title_run.bold = True
    title_run.font.size = Pt(22)
    title_run.font.color.rgb = RGBColor(37, 99, 235)

    document.add_paragraph(f"综合评分：{result_summary['overall_score']}")
    document.add_paragraph(f"总体评价：{summary}")

    document.add_heading("五维评分详情", level=2)
    table = document.add_table(rows=1, cols=4)
    table.style = "Table Grid"
    header_cells = table.rows[0].cells
    for index, header in enumerate(("维度", "分数", "分析", "建议")):
        header_cells[index].text = header
    for item in dimensions:
        if not isinstance(item, dict):
            continue
        row = table.add_row().cells
        row[0].text = _dimension_label(str(item.get("name", "")))
        row[1].text = str(item.get("score", ""))
        row[2].text = str(item.get("analysis", ""))
        row[3].text = str(item.get("suggestion", ""))

    warnings = result_summary.get("artifact_warnings", [])
    if isinstance(warnings, list) and warnings:
        document.add_heading("产物告警", level=2)
        for warning in warnings:
            if isinstance(warning, dict):
                document.add_paragraph(str(warning.get("message", "")), style="List Bullet")

    if radar_chart:
        document.add_heading("综合能力雷达图", level=2)
        try:
            document.add_picture(BytesIO(radar_chart), width=Inches(5.2))
            picture_paragraph = document.paragraphs[-1]
            picture_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        except Exception:
            document.add_paragraph("雷达图暂不可嵌入，评分结果仍可查看。")

    output = BytesIO()
    document.save(output)
    return output.getvalue()


def _score(item: dict[str, Any]) -> int:
    value = item.get("score")
    if not isinstance(value, int) or value < 0 or value > 100:
        raise ValueError("dimension score is invalid")
    return value


def _dimension_label(name: str) -> str:
    labels = {
        "technology_stack": "技术栈匹配",
        "experience": "经验匹配",
        "education": "学历背景",
        "potential": "发展潜力",
        "resume_structure": "简历结构",
    }
    return labels.get(name, name)
